from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "DOCUMENTACION_FLUJO_LIA.md"
OUTPUT = ROOT / "DOCUMENTACION_FLUJO_LIA.docx"

NAVY = "17365D"
TEAL = "0F6B78"
LIGHT_TEAL = "DDEBF1"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"
BLACK = "1F2937"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, *, bold=False, color=BLACK, size=8.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Actualizar campo"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])


def add_page_number(paragraph) -> None:
    paragraph.add_run("Página ")
    add_field(paragraph, "PAGE")
    paragraph.add_run(" de ")
    add_field(paragraph, "NUMPAGES")


INLINE_RE = re.compile(
    r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)"
)


def add_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(TEAL)
        elif token.startswith("["):
            label, target = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label} ({target})")
            run.font.color.rgb = RGBColor.from_string(TEAL)
            run.underline = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    heading_specs = {
        "Title": (26, NAVY, 12, 10),
        "Subtitle": (12, TEAL, 3, 6),
        "Heading 1": (17, NAVY, 14, 7),
        "Heading 2": (13.5, TEAL, 11, 5),
        "Heading 3": (11.5, NAVY, 8, 4),
        "Heading 4": (10.5, TEAL, 6, 3),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(3)

    if "Technical Block" not in styles:
        technical = styles.add_style("Technical Block", WD_STYLE_TYPE.PARAGRAPH)
        technical.font.name = "Consolas"
        technical.font.size = Pt(8)
        technical.font.color.rgb = RGBColor.from_string(BLACK)
        technical.paragraph_format.left_indent = Cm(0.4)
        technical.paragraph_format.right_indent = Cm(0.4)
        technical.paragraph_format.space_before = Pt(3)
        technical.paragraph_format.space_after = Pt(3)

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        callout.font.name = "Aptos"
        callout.font.size = Pt(9.5)
        callout.font.italic = True
        callout.font.color.rgb = RGBColor.from_string(NAVY)
        callout.paragraph_format.left_indent = Cm(0.6)
        callout.paragraph_format.right_indent = Cm(0.4)
        callout.paragraph_format.space_before = Pt(5)
        callout.paragraph_format.space_after = Pt(8)


def add_paragraph_shading(paragraph, fill: str, border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "7")
        left.set(qn("w:color"), border)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)

    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = label.add_run("DOCUMENTACIÓN TÉCNICA Y FUNCIONAL")
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("LIA")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(38)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Flujo integral de gestión de incapacidades")
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run("━" * 34)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    statement = doc.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement.paragraph_format.left_indent = Cm(2.2)
    statement.paragraph_format.right_indent = Cm(2.2)
    run = statement.add_run(
        "Arquitectura central 100 % in-house para orquestación, APIs, "
        "persistencia, almacenamiento documental, autenticación y trazabilidad."
    )
    run.font.name = "Aptos"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(BLACK)

    doc.add_paragraph()
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    values = [
        ("Versión documental", "1.0"),
        ("Fecha", "30 de julio de 2026"),
        ("Propietario sugerido", "Gestión Humana / Tecnología"),
        ("Clasificación", "Uso interno"),
    ]
    for row, (key, value) in zip(meta.rows, values):
        row.cells[0].width = Cm(5.5)
        row.cells[1].width = Cm(7)
        set_cell_text(row.cells[0], key, bold=True, color=NAVY, size=9)
        set_cell_text(row.cells[1], value, color=BLACK, size=9)
        set_cell_shading(row.cells[0], LIGHT_TEAL)

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(65)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("GRUPO EMPRESARIAL EN LÍNEA S.A.")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    doc.add_page_break()


def add_document_control(doc: Document) -> None:
    doc.add_heading("Control documental", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Versión", "Fecha", "Descripción", "Estado"]
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE, size=8.5)
        set_cell_shading(table.rows[0].cells[index], NAVY)
    values = ["1.0", "30/07/2026", "Emisión inicial de documentación integral", "Para revisión"]
    row = table.add_row()
    for index, value in enumerate(values):
        set_cell_text(row.cells[index], value, size=8.5)
        if index % 2 == 0:
            set_cell_shading(row.cells[index], LIGHT_GRAY)
    set_repeat_table_header(table.rows[0])

    doc.add_heading("Tabla de contenido", level=1)
    note = doc.add_paragraph()
    note.add_run(
        "En Microsoft Word, haga clic derecho sobre la tabla y seleccione "
        "«Actualizar campo» para recalcular títulos y páginas."
    ).italic = True
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(cells)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for column in range(column_count):
        value = rows[0][column] if column < len(rows[0]) else ""
        set_cell_text(table.rows[0].cells[column], value, bold=True, color=WHITE, size=8)
        set_cell_shading(table.rows[0].cells[column], NAVY)
    set_repeat_table_header(table.rows[0])
    for row_index, source_row in enumerate(rows[1:], start=1):
        target = table.add_row()
        for column in range(column_count):
            value = source_row[column] if column < len(source_row) else ""
            value = re.sub(r"\*\*([^*]+)\*\*", r"\1", value)
            value = re.sub(r"`([^`]+)`", r"\1", value)
            set_cell_text(target.cells[column], value, size=7.8)
            if row_index % 2 == 0:
                set_cell_shading(target.cells[column], LIGHT_GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, lines: list[str], language: str) -> None:
    label = "DIAGRAMA — CÓDIGO MERMAID" if language == "mermaid" else "BLOQUE TÉCNICO"
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(0)
    run = title.add_run(label)
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(WHITE)
    add_paragraph_shading(title, TEAL)

    body = doc.add_paragraph(style="Technical Block")
    body.paragraph_format.keep_together = True
    body.add_run("\n".join(lines))
    add_paragraph_shading(body, LIGHT_GRAY, TEAL)


def add_body_from_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    # The Word cover replaces the Markdown H1 and its metadata table.
    index = 0
    if lines and lines[0].startswith("# "):
        index = 1
    if index < len(lines) and not lines[index].strip():
        index += 1
    if index < len(lines) and lines[index].strip().startswith("| Campo |"):
        _, index = parse_table(lines, index)

    in_code = False
    code_language = ""
    code_lines: list[str] = []
    skip_approval_heading = False

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_language = stripped[3:].strip().lower()
                code_lines = []
            else:
                add_code_block(doc, code_lines, code_language)
                in_code = False
                code_language = ""
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{2,6})\s+(.+)$", raw)
        if heading:
            level = min(len(heading.group(1)) - 1, 4)
            text = heading.group(2)
            if text == "Aprobación documental":
                skip_approval_heading = True
            doc.add_heading(text, level=level)
            index += 1
            continue

        if stripped == "---":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run("━" * 72)
            run.font.color.rgb = RGBColor.from_string(TEAL)
            index += 1
            continue

        if stripped.startswith(">"):
            paragraph = doc.add_paragraph(style="Callout")
            add_inline(paragraph, stripped.lstrip("> ").strip())
            add_paragraph_shading(paragraph, LIGHT_BLUE, TEAL)
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if numbered:
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(2))
            index += 1
            continue
        if bullet:
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(1))
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1

    if in_code:
        add_code_block(doc, code_lines, code_language)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(1.8)
        section.header_distance = Cm(0.7)
        section.footer_distance = Cm(0.7)

        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("LIA  |  Gestión de incapacidades")
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)

        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.style = doc.styles["Normal"]
        for run in paragraph.runs:
            run.font.size = Pt(8)
        add_page_number(paragraph)


def set_document_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "LIA — Documentación integral del flujo de incapacidades"
    props.subject = "Arquitectura, operación, seguridad y soporte de la solución LIA"
    props.author = "GRUPO EMPRESARIAL EN LÍNEA S.A."
    props.keywords = "LIA, incapacidades, n8n, in-house, PostgreSQL, MinIO"
    props.comments = "Generado a partir de DOCUMENTACION_FLUJO_LIA.md"

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"No se encontró la fuente: {SOURCE}")

    markdown = SOURCE.read_text(encoding="utf-8")
    document = Document()
    configure_styles(document)
    configure_sections(document)
    set_document_properties(document)
    add_cover(document)
    add_document_control(document)
    add_body_from_markdown(document, markdown)
    document.save(OUTPUT)
    print(f"Word generado: {OUTPUT}")
    print(f"Tamaño: {OUTPUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
